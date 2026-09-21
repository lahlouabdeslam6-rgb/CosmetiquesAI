"""Editable Word exports built from the persisted client dossier and selected formula."""
from io import BytesIO
from typing import Any, Literal

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def _text(value: Any) -> str:
    if isinstance(value, list):
        return ", ".join(map(str, value)) or "Non renseigné"
    if isinstance(value, bool):
        return "Oui" if value else "Non"
    return str(value) if value is not None and str(value).strip() else "Non renseigné"


def _table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    for column, width in zip(table.columns, widths):
        column.width = Inches(width)
    for cell, title, width in zip(table.rows[0].cells, headers, widths):
        cell.width = Inches(width)
        run = cell.paragraphs[0].add_run(title)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        shade = OxmlElement("w:shd"); shade.set(qn("w:fill"), "333333"); cell._tc.get_or_add_tcPr().append(shade)
    repeat = OxmlElement("w:tblHeader"); table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for index, values in enumerate(rows):
        cells = table.add_row().cells
        for cell, value, width in zip(cells, values, widths):
            cell.width = Inches(width); cell.text = _text(value)
            if index % 2 == 0:
                shade = OxmlElement("w:shd"); shade.set(qn("w:fill"), "F4F4F4"); cell._tc.get_or_add_tcPr().append(shade)
    for row in table.rows:
        for cell in row.cells:
            props = cell._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            for edge in ("top", "left", "bottom", "right"):
                element = OxmlElement(f"w:{edge}"); element.set(qn("w:val"), "single"); element.set(qn("w:sz"), "4"); element.set(qn("w:color"), "D9D9D9"); borders.append(element)
            props.append(borders)
            margins = OxmlElement("w:tcMar")
            for edge in ("top", "left", "bottom", "right"):
                element = OxmlElement(f"w:{edge}"); element.set(qn("w:w"), "110"); element.set(qn("w:type"), "dxa"); margins.append(element)
            props.append(margins)
    doc.add_paragraph()


def build_formulation_document(dossier: dict[str, Any], kind: Literal["cahier", "procedure"]) -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5); section.page_height = Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.75)
    section.left_margin = section.right_margin = Inches(0.8)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"; normal.font.size = Pt(11); normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.1
    for name in ("Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[name]; style.font.color.rgb = RGBColor(0, 0, 0)
    section.header.paragraphs[0].text = "MIELMAR    Laboratoire"
    section.header.paragraphs[0].style = "Caption"
    footer = section.footer.paragraphs[0]
    footer.add_run(f"{dossier['reference']}    Page ")
    field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE"); footer._p.append(field)
    doc.add_heading("Cahier des charges client" if kind == "cahier" else "Procédé de fabrication laboratoire", 0)
    doc.add_paragraph("Ce document reprend les informations enregistrées du dossier client." if kind == "cahier" else "Ce document décrit la formule sélectionnée et son procédé proposé pour préparer les essais au laboratoire. Les paramètres restent à confirmer au lot pilote et à valider par le formulateur avant fabrication.")
    for label, value in (("Référence", dossier["reference"]), ("Client", dossier["client_name"]), ("Projet", dossier.get("project_name")), ("Produit souhaité", dossier["product_type"]), ("Version", dossier.get("revision", 0)), ("Statut", dossier["status"])):
        paragraph = doc.add_paragraph(); paragraph.add_run(f"{label} : ").bold = True; paragraph.add_run(_text(value))
    if kind == "cahier":
        doc.add_heading("Contact du client", 1)
        doc.add_paragraph(f"Email : {_text(dossier.get('email'))}\nTéléphone : {_text(dossier.get('phone'))}")
        doc.add_heading("Besoin et résultat attendu", 1)
        doc.add_paragraph(_text(dossier.get("needs")))
        doc.add_heading("Réponses au questionnaire", 1)
        for answer in dossier.get("answer_details", []):
            doc.add_heading(answer["label"], 2)
            doc.add_paragraph(_text(answer.get("value")))
    else:
        formula = dossier.get("formula") or {}
        report = dossier.get("lab_report") or {}
        doc.add_heading("Produit et lot pilote", 1)
        doc.add_paragraph(f"{formula.get('suggested_name') or dossier['product_type']}\nLot pilote : {formula.get('batch_size_g', '—')} g\nTotal : {formula.get('total_percentage', '—')} % m/m\nProposition : {dossier.get('selected_formula_index', 0) + 1}")
        doc.add_paragraph(_text(report.get("summary")))
        doc.add_heading("Formule quantitative", 1)
        _table(doc, ["Matière", "% m/m", "Quantité", "Rôle"], [(item.get("name"), item.get("percentage"), f"{item.get('required_quantity', round(float(item.get('percentage', 0)) * float(formula.get('batch_size_g', 0)) / 100, 4))} {item.get('unit', 'g')}", item.get("role")) for item in formula.get("ingredients", [])], [2.2, 0.7, 1.1, 2.9])
        doc.add_heading("Mode opératoire proposé", 1)
        steps = report.get("manufacturing_steps") or []
        if steps:
            for index, step in enumerate(steps, 1):
                doc.add_heading(f"Étape {index} {step['phase']} {step['title']}", 2)
                for label, key in (("Matières", "materials"), ("Opération", "instruction"), ("Température", "temperature"), ("Durée et attente", "duration"), ("Agitation", "mixing"), ("Critère de fin et contrôle", "checkpoint")):
                    paragraph = doc.add_paragraph(); paragraph.add_run(f"{label} : ").bold = True; paragraph.add_run(step[key])
        else:
            doc.add_paragraph("Le procédé détaillé n'a pas encore été généré. Les indications existantes ci-dessous ne constituent pas un mode opératoire complet.")
            for step in report.get("manufacturing_process", []): doc.add_paragraph(str(step), "List Number")
        for title, key in (("Paramètres à confirmer avant essais", "process_parameters_to_confirm"), ("Tests de stabilité", "stability_tests"), ("Tests et preuves d’efficacité", "efficacy_tests"), ("Limites et validations", "warnings")):
            values = report.get(key) or []
            if values:
                doc.add_heading(title, 1)
                for value in values: doc.add_paragraph(str(value), "List Bullet")
        if dossier.get("missing_materials"):
            doc.add_heading("Matières à approvisionner", 1)
            for item in dossier["missing_materials"]: doc.add_paragraph(f"{item['name']} — {item['required_quantity']} {item['unit']}", "List Bullet")
        doc.add_heading("Validation administrateur", 1)
        doc.add_paragraph(f"Formule : {'validée' if dossier['status'] == 'approved' else 'en attente de validation'}.\nProduit enregistré : {_text(dossier.get('product_id'))}.\nValidation : {_text(dossier.get('approved_at'))}.")
    output = BytesIO(); doc.save(output); return output.getvalue()
