import json

import pytest
from sqlalchemy import select, func
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from app.api.cosmetic_formulation import (
    CosmeticRequestWrite,
    MaterialPatch,
    MaterialWrite,
    create_cosmetic_request,
    create_raw_material,
    generate_formula,
    list_raw_materials,
    update_raw_material,
)
from app.db.sqlite import Base
from app.models.cosmetic_formulation import CosmeticRequest, RawMaterial
from app.models.operations import AuditLog, Notification
from app.models.site_setting import SiteSetting
from app.models.user import User
from app.services.laboratory_agent import validate_proposal
from app.services.laboratory_agent import (
    assemble_manufacturing_plan,
    build_local_manufacturing_plan,
    manufacturing_plan_schema,
    validate_manufacturing_steps,
)
from app.services.formulation_documents import build_formulation_document
from io import BytesIO
from docx import Document


def detailed_process():
    return {"steps": [{"phase": f"Phase {index}", "title": "Préparation", "materials": "Aqua 96 g et Glycérine 4 g", "instruction": "Incorporer progressivement et vérifier la dispersion.", "temperature": "À confirmer selon fiche fournisseur", "duration": "10 minutes proposées à confirmer au pilote", "mixing": "Agitation douce adaptée au matériel", "checkpoint": "Mélange homogène sans particules"} for index in range(1, 7)], "unresolved_parameters": ["Confirmer les paramètres au lot pilote"]}


def test_manufacturing_procedure_requires_every_operating_parameter():
    payload = detailed_process()
    result = validate_manufacturing_steps(payload)
    assert len(result["manufacturing_steps"]) == 6
    assert "10 minutes" in result["manufacturing_process"][0]
    payload["steps"][0]["temperature"] = ""
    with pytest.raises(ValueError, match="incomplète"):
        validate_manufacturing_steps(payload)


def test_immediate_procedure_is_complete_and_grounded_in_the_client_dossier():
    formula = {
        "batch_size_g": 250,
        "ingredients": [
            {"name": "Aqua", "percentage": 89.8, "role": "Solvant"},
            {"name": "Cetyl Palmitate", "percentage": 5, "role": "Agent de texture"},
            {"name": "Niacinamide", "percentage": 4, "role": "Actif"},
            {"name": "Benzyl Alcohol", "percentage": 1.2, "role": "Conservateur"},
        ],
    }
    brief = {
        "answers": {
            "target_ph": "5,2–5,8",
            "viscosity": "20 000–30 000 mPa.s",
            "packaging_type": "Flacon airless",
        }
    }

    result = build_local_manufacturing_plan(formula, brief)

    assert result["process_status"] == "proposed"
    assert result["process_source"] == "laboratory_agent_calculated_immediate"
    assert len(result["manufacturing_steps"]) == 8
    procedure = " ".join(
        str(value)
        for step in result["manufacturing_steps"]
        for value in step.values()
    )
    for ingredient in formula["ingredients"]:
        assert ingredient["name"] in procedure
    assert "250 g" in procedure
    assert "5,2–5,8" in procedure
    assert "20 000–30 000 mPa.s" in procedure
    assert "Flacon airless" in procedure


def test_procedure_checks_ingredient_coverage_weighing_and_client_controls():
    payload = detailed_process()
    formula = {"ingredients": [{"name": "Aqua"}, {"name": "Glycérine"}]}
    brief = {"answers": {"target_ph": "Cible client", "viscosity": "Cible client"}}
    with pytest.raises(ValueError, match="pesées"):
        validate_manufacturing_steps(payload, formula, brief)
    payload["steps"][0]["title"] = "Pesées séparées"
    with pytest.raises(ValueError, match="pH"):
        validate_manufacturing_steps(payload, formula, brief)
    payload["steps"][4]["instruction"] += " Contrôler le pH et la viscosité avant de poursuivre."
    assert validate_manufacturing_steps(payload, formula, brief)["process_status"] == "proposed"
    formula["ingredients"].append({"name": "Actif oublié"})
    with pytest.raises(ValueError, match="Actif oublié"):
        validate_manufacturing_steps(payload, formula, brief)


def test_process_assignments_keep_every_ingredient_quantity_and_client_control():
    formula = {"batch_size_g": 250, "ingredients": [{"name": "Aqua", "percentage": 96}, {"name": "Glycérine", "percentage": 4}]}
    payload = detailed_process()
    payload["ingredient_steps"] = {"Aqua": 1, "Glycérine": 2}
    payload["steps"][0]["instruction"] = "Incorporer Aqua progressivement."
    payload["steps"][1]["instruction"] = "Incorporer Glycérine progressivement."
    brief = {"answers": {"target_ph": "5,3 à 5,8", "viscosity": "Cible client à 25 °C"}}
    schema = manufacturing_plan_schema(formula)
    assert schema["properties"]["ingredient_steps"]["required"] == ["Aqua", "Glycérine"]
    result = assemble_manufacturing_plan(payload, formula, brief)
    steps = result["manufacturing_steps"]
    assert steps[1]["title"].startswith("Peser")
    assert "Aqua : 240 g" in steps[1]["materials"]
    assert "Glycérine : 10 g" in steps[3]["materials"]
    assert any("5,3 à 5,8" in step["instruction"] for step in steps)
    del payload["ingredient_steps"]["Glycérine"]
    with pytest.raises(ValueError, match="Glycérine"):
        assemble_manufacturing_plan(payload, formula, brief)


def test_word_exports_preserve_client_brief_and_selected_procedure():
    dossier = {"reference": "COS-TEST", "client_name": "Karim", "product_type": "Gel", "status": "received", "needs": "Hydratant", "email": "client@example.test", "phone": "0600000000", "answer_details": [{"label": "Packaging", "value": ["Airless"]}], "formula": {"suggested_name": "Gel Karim", "batch_size_g": 100, "total_percentage": 100, "ingredients": [{"name": "Aqua", "percentage": 96, "required_quantity": 96, "unit": "g", "role": "Solvant"}, {"name": "Glycérine", "percentage": 4, "required_quantity": 4, "unit": "g", "role": "Humectant"}]}, "lab_report": {"summary": "Gel hydratant", **validate_manufacturing_steps(detailed_process())}}
    brief = Document(BytesIO(build_formulation_document(dossier, "cahier")))
    text = "\n".join(p.text for p in brief.paragraphs)
    assert "Karim" in text and "Airless" in text and "Hydratant" in text
    procedure = Document(BytesIO(build_formulation_document(dossier, "procedure")))
    text = "\n".join(p.text for p in procedure.paragraphs)
    assert "Durée et attente" in text and "10 minutes" in text and "Phase 6" in text
    assert procedure.tables[0].cell(1, 0).text == "Aqua"


@pytest.mark.asyncio
async def test_formula_approval_records_client_after_approved_procurement(monkeypatch):
    from fastapi import BackgroundTasks, HTTPException
    from app.api.cosmetic_formulation import approve_formula, decide_missing_materials, FormulaApprovalWrite, MaterialDecisionWrite
    from app.schemas.product import ProductResponse
    engine = create_async_engine("sqlite+aiosqlite:///:memory:")
    async with engine.begin() as connection:
        await connection.run_sync(Base.metadata.create_all)
    Session = async_sessionmaker(engine, expire_on_commit=False)
    created = []
    async def fake_create(self, data):
        created.append(data.model_dump())
        return ProductResponse(product_id=42, name=data.name, stock=data.stock, is_active=data.is_active)
    monkeypatch.setattr("app.services.product.ProductService.create_product", fake_create)
    async with Session() as session:
        admin = User(email="admin@example.test", role="admin", hashed_password="test")
        customer = User(email="client@example.test", full_name="Karim", role="customer", hashed_password="test")
        session.add_all([admin, customer]); await session.commit()
        row = CosmeticRequest(reference="COS-TEST", user_id=customer.id, client_name="Karim", email=customer.email, phone="0600000000", product_type="Gel", status="needs_materials", formula_json=json.dumps({"total_percentage": 100, "ingredients": [{"name": "Aqua", "percentage": 96}, {"name": "Glycerin", "percentage": 4}]}), lab_report_json=json.dumps({"summary": "Gel hydratant"}), missing_materials_json=json.dumps([{"name": "Glycérine", "required_quantity": 4, "unit": "g"}]))
        session.add(row); await session.commit()
        data = FormulaApprovalWrite(product_name="Gel Karim", media_mode="later")
        with pytest.raises(HTTPException) as error:
            await approve_formula(row.id, data, BackgroundTasks(), admin, session)
        assert error.value.status_code == 409
        await decide_missing_materials(row.id, MaterialDecisionWrite(can_source=True), admin, session)
        row.lab_report_json = json.dumps({"process_status": "failed"})
        await session.commit()
        with pytest.raises(HTTPException) as failed_process:
            await approve_formula(row.id, data, BackgroundTasks(), admin, session)
        assert failed_process.value.status_code == 409 and not created
        row.lab_report_json = json.dumps({
            "summary": "Gel hydratant",
            "process_status": "proposed",
            "manufacturing_steps": [{
                "phase": "Phase finale",
                "title": "Conditionnement",
                "materials": "Mélange final",
                "instruction": "Conditionner le lot pilote contrôlé.",
                "temperature": "Température ambiante",
                "duration": "Après libération des contrôles",
                "mixing": "Sans agitation",
                "checkpoint": "Lot identifié et conforme",
            }],
        })
        await session.commit()
        result = await approve_formula(row.id, data, BackgroundTasks(), admin, session)
        assert result["status"] == "approved" and result["product_id"] == 42
        assert created[0]["custom_client_name"] == "Karim"
        assert created[0]["custom_client_user_id"] == customer.id
        assert created[0]["cosmetic_request_reference"] == "COS-TEST"
        assert created[0]["stock"] == 0 and created[0]["is_active"] is False
        await approve_formula(row.id, data, BackgroundTasks(), admin, session)
        assert len(created) == 1
    await engine.dispose()


def test_laboratory_agent_completes_formula_with_vehicle_to_100_percent():
    aqua = RawMaterial(id=1, name="Aqua", quantity=1000, reserved_quantity=0, unit="g", catalog_json="{}")
    payload = validate_proposal({
        "suggested_name": "Soin libre", "product_concept": "Concept client", "batch_size_g": 250,
        "ingredients": [
            {"material_id": None, "name": "Actif proposé", "percentage": 3, "role": "Actif", "justification": "Objectif client", "source": "purchase"},
            {"material_id": None, "name": "Conservateur proposé", "percentage": 2, "role": "Conservation", "justification": "Protège le prototype", "source": "purchase"},
        ],
    }, [aqua], {"product_idea": "Soin visage aqueux"})
    assert sum(item["percentage"] for item in payload["ingredients"]) == 100
    assert payload["ingredients"][0]["name"] == "Aqua"


def test_laboratory_agent_resolves_null_qsp_and_percentage_strings():
    aqua = RawMaterial(id=1, name="Aqua", quantity=1000, reserved_quantity=0, unit="g", catalog_json="{}")
    payload = validate_proposal({
        "suggested_name": "Gel hydratant", "product_concept": "Gel visage", "batch_size_g": 500,
        "ingredients": [
            {"material_id": 1, "name": "Aqua", "percentage": None, "role": "Solvant q.s.p. 100 %", "justification": "Phase aqueuse", "source": "stock"},
            {"material_id": None, "name": "Glycérine", "percentage": "4,0 %", "role": "Humectant", "justification": "Hydratation", "source": "purchase"},
            {"material_id": None, "name": "Gélifiant", "percentage": 1, "role": "Gélifiant", "justification": "Texture", "source": "purchase"},
        ],
    }, [aqua], {"product_idea": "Gel visage aqueux"})
    assert sum(item["percentage"] for item in payload["ingredients"]) == 100
    assert payload["ingredients"][0]["name"] == "Aqua"
    assert payload["ingredients"][0]["percentage"] == 95
    assert payload["ingredients"][1]["percentage"] == 4


@pytest.mark.asyncio
async def test_open_product_brief_is_formulated_then_checked_against_stock(monkeypatch):
    engine = create_async_engine("sqlite+aiosqlite:///:memory:")
    async with engine.begin() as connection:
        await connection.run_sync(Base.metadata.create_all)
    Session = async_sessionmaker(engine, expire_on_commit=False)
    async with Session() as session:
        admin = User(email="admin@example.test", role="admin", hashed_password="test-only")
        customer = User(email="client@example.test", full_name="Client Test", role="customer", hashed_password="test-only")
        session.add_all([admin, customer]); await session.commit(); await session.refresh(admin); await session.refresh(customer)
        material = await create_raw_material(
            MaterialWrite(name="Matière test", unit="g", quantity=2, minimum_quantity=1),
            admin, session,
        )
        assert material["available_quantity"] == 2
        request_result = await create_cosmetic_request(
            CosmeticRequestWrite(project_name="Projet test", phone="0600000000", product_type="Soin test", needs="Besoin test", answers={"tests":["Stabilité","Hydratation / Cornéométrie"]}),
            customer, session,
        )
        request = await session.scalar(select(CosmeticRequest).where(CosmeticRequest.reference == request_result["reference"]))
        async def fake_generate(_self, **kwargs):
            assert kwargs["product_idea"] == "Soin test"
            assert kwargs["needs"] == "Besoin test"
            return {
                "suggested_name": "Soin test unique", "product_concept": "Concept créé depuis le cahier des charges.",
                "batch_size_g": 1000, "ingredients": [{"material_id": material["id"], "name": "Matière test", "percentage": 100, "role": "Base hydratante", "justification": "Répond au besoin client.", "source": "stock", "material_catalog": {}}],
                "manufacturing_process": ["Peser", "Mélanger"], "stability_tests": ["Stabilité"],
                "efficacy_tests": ["Hydratation / Cornéométrie"], "consistency_and_additives": [],
                "hydrating_and_repair_actives": ["Matière test"], "warnings": [],
            }
        monkeypatch.setattr("app.services.laboratory_agent.LaboratoryAgentService.generate", fake_generate)
        from fastapi import BackgroundTasks
        automatic_tasks = BackgroundTasks()
        missing = await generate_formula(request.id, admin, session, background_tasks=automatic_tasks)
        assert len(automatic_tasks.tasks) == 1
        assert missing["lab_report"]["process_status"] == "generating"
        assert missing["status"] == "needs_materials"
        assert missing["missing_materials"][0]["missing"] == 998
        await update_raw_material(material["id"], MaterialPatch(quantity=1200), admin, session)
        ready = await generate_formula(request.id, admin, session)
        assert ready["status"] == "draft_ready"
        assert ready["formula"]["total_percentage"] == 100
        assert ready["formula"]["ingredients"][0]["required_quantity"] == 1000
        assert ready["lab_report"]["manufacturing_process"]
        assert ready["lab_report"]["stability_tests"]
        assert ready["lab_report"]["efficacy_tests"] == ["Hydratation / Cornéométrie"]
        assert ready["user_id"] == customer.id
        assert ready["client_snapshot"]["email"] == customer.email
        stock = await list_raw_materials(admin, session, False)
        assert stock["total"] == 1 and stock["low_stock"] == 0
        assert await session.scalar(select(func.count()).select_from(Notification)) == 3
        notification_types = set((await session.scalars(select(Notification.notification_type))).all())
        assert notification_types == {"cosmetic_request", "raw_materials_missing", "formula_ready"}
        assert await session.scalar(select(func.count()).select_from(AuditLog)) == 4
    await engine.dispose()


@pytest.mark.asyncio
async def test_reserved_material_cannot_exceed_quantity():
    engine = create_async_engine("sqlite+aiosqlite:///:memory:")
    async with engine.begin() as connection:
        await connection.run_sync(Base.metadata.create_all)
    Session = async_sessionmaker(engine, expire_on_commit=False)
    async with Session() as session:
        admin = User(email="admin@example.test", role="admin", hashed_password="test-only")
        session.add(admin); await session.commit(); await session.refresh(admin)
        material = await create_raw_material(MaterialWrite(name="Matière test", quantity=2), admin, session)
        from fastapi import HTTPException
        with pytest.raises(HTTPException) as exc:
            await update_raw_material(material["id"], MaterialPatch(reserved_quantity=3), admin, session)
        assert exc.value.status_code == 422
    await engine.dispose()


@pytest.mark.asyncio
async def test_procedure_job_updates_selected_variant_and_discards_stale_result(monkeypatch):
    from fastapi import BackgroundTasks
    from app.api.cosmetic_formulation import generate_manufacturing_procedure, run_manufacturing_procedure, download_formulation_word
    from fastapi import HTTPException
    engine = create_async_engine("sqlite+aiosqlite:///:memory:")
    async with engine.begin() as connection:
        await connection.run_sync(Base.metadata.create_all)
    Session = async_sessionmaker(engine, expire_on_commit=False)
    monkeypatch.setattr("app.api.cosmetic_formulation.SqliteSession", Session)
    async with Session() as session:
        admin = User(email="admin@example.test", role="admin", hashed_password="test")
        session.add(admin); await session.commit()
        formula = {"batch_size_g": 100, "total_percentage": 100, "ingredients": [{"name": "Aqua", "percentage": 96}, {"name": "Glycerin", "percentage": 4}]}
        variant = {"formula": formula, "lab_report": {}, "missing_materials": [], "status": "draft_ready", "version": 1}
        row = CosmeticRequest(reference="COS-JOB", user_id=admin.id, client_name="Karim", email=admin.email, phone="0600000000", product_type="Gel", status="draft_ready", formula_json=json.dumps(formula), formulas_json=json.dumps([variant]), lab_report_json="{}")
        session.add(row); await session.commit()
        with pytest.raises(HTTPException) as incomplete:
            await download_formulation_word(row.id, admin, session, "procedure")
        assert incomplete.value.status_code == 409
        tasks = BackgroundTasks()
        queued = await generate_manufacturing_procedure(row.id, tasks, admin, session)
        assert queued["lab_report"]["process_status"] == "generating"
        assert len(tasks.tasks) == 1
        await tasks()
        await session.refresh(row)
        saved = json.loads(row.lab_report_json)
        assert saved["process_status"] == "proposed"
        assert len(saved["manufacturing_steps"]) == 8
        assert saved["process_source"] == "laboratory_agent_calculated_immediate"
        exported = await download_formulation_word(row.id, admin, session, "procedure")
        assert exported.media_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        assert Document(BytesIO(exported.body)).tables[0].cell(1, 0).text == "Aqua"
        assert json.loads(row.formulas_json)[0]["lab_report"] == saved
        await run_manufacturing_procedure(row.id, 0, "outdated-job")
        await session.refresh(row)
        assert json.loads(row.lab_report_json) == saved
        # A different proposal receives its own procedure without replacing the selection.
        variants = json.loads(row.formulas_json)
        variants.append({**variant, "lab_report": {"process_status": "generating", "process_job_id": "second-proposal"}})
        row.formulas_json = json.dumps(variants)
        await session.commit()
        await run_manufacturing_procedure(row.id, 1, "second-proposal")
        await session.refresh(row)
        assert json.loads(row.lab_report_json) == saved
        assert json.loads(row.formulas_json)[1]["lab_report"]["process_status"] == "proposed"
    await engine.dispose()


def test_procedure_rejects_placeholder_parameters_and_false_ingredient_assignment():
    payload = detailed_process()
    payload["steps"][0]["temperature"] = "-"
    with pytest.raises(ValueError, match="incomplète"):
        validate_manufacturing_steps(payload)
    payload = detailed_process()
    payload["ingredient_steps"] = {"Aqua": 1, "Glycérine": 2}
    formula = {"batch_size_g": 100, "ingredients": [{"name": "Aqua", "percentage": 96}, {"name": "Glycérine", "percentage": 4}]}
    with pytest.raises(ValueError, match="Affectation incohérente"):
        assemble_manufacturing_plan(payload, formula, {})
